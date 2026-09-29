"""Search-task definitions and spreadsheet-backed product input."""

from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from pathlib import Path

from .config import ConfigurationError
from .models import clean_text, normalize_barcode


PRODUCT_SHEET = "KVI_100"
PRODUCT_NAME_HEADER = "מוצר / SKU"
PRODUCT_BARCODE_HEADER = "EAN / ברקוד"


class SearchKind(str, Enum):
    KEYWORD = "keyword"
    PRODUCT = "product"


class ResultPolicy(str, Enum):
    ALL = "all"
    FIRST = "first"


@dataclass(frozen=True)
class SearchTask:
    query: str
    checkpoint_key: str
    kind: SearchKind
    result_policy: ResultPolicy
    expected_barcode: str = ""


def keyword_checkpoint_key(keyword: str) -> str:
    return f"keyword:{keyword}"


def product_checkpoint_key(barcode: str) -> str:
    return f"product:{barcode}"


def keyword_search_tasks(keywords: list[str]) -> list[SearchTask]:
    return [
        SearchTask(
            query=keyword,
            checkpoint_key=keyword_checkpoint_key(keyword),
            kind=SearchKind.KEYWORD,
            result_policy=ResultPolicy.ALL,
        )
        for keyword in keywords
    ]


def _barcode_cell_text(value: object) -> str:
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(str(value)) if value is not None else ""


def load_product_searches(path: Path) -> list[SearchTask]:
    """Load ordered product-name searches from the workbook's KVI sheet."""

    if not path.exists():
        raise ConfigurationError(f"Product workbook does not exist: {path}")

    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ConfigurationError(
            "Product workbook support requires openpyxl; install requirements.txt."
        ) from error

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ConfigurationError(f"Could not read product workbook {path}: {error}") from error

    try:
        if PRODUCT_SHEET not in workbook.sheetnames:
            raise ConfigurationError(
                f"Product workbook {path} does not contain sheet {PRODUCT_SHEET!r}"
            )
        worksheet = workbook[PRODUCT_SHEET]
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as error:
            raise ConfigurationError(
                f"Product sheet {PRODUCT_SHEET!r} is empty in {path}"
            ) from error

        headers = {
            clean_text(str(value)): index
            for index, value in enumerate(raw_headers)
            if value is not None and clean_text(str(value))
        }
        missing_headers = [
            header
            for header in (PRODUCT_NAME_HEADER, PRODUCT_BARCODE_HEADER)
            if header not in headers
        ]
        if missing_headers:
            raise ConfigurationError(
                f"Product sheet {PRODUCT_SHEET!r} is missing required columns: "
                + ", ".join(missing_headers)
            )

        name_index = headers[PRODUCT_NAME_HEADER]
        barcode_index = headers[PRODUCT_BARCODE_HEADER]
        tasks: list[SearchTask] = []
        seen_barcodes: set[str] = set()
        for row_number, row in enumerate(rows, start=2):
            raw_name = row[name_index] if name_index < len(row) else None
            raw_barcode = row[barcode_index] if barcode_index < len(row) else None
            name = clean_text(str(raw_name)) if raw_name is not None else ""
            barcode_text = _barcode_cell_text(raw_barcode)
            if not name and not barcode_text:
                continue
            if not name:
                raise ConfigurationError(
                    f"Missing product name in {PRODUCT_SHEET!r} row {row_number}"
                )
            barcode = normalize_barcode(barcode_text)
            if not barcode:
                raise ConfigurationError(
                    f"Invalid product barcode in {PRODUCT_SHEET!r} row {row_number}: "
                    f"{barcode_text!r}"
                )
            if barcode in seen_barcodes:
                raise ConfigurationError(
                    f"Duplicate product barcode in {PRODUCT_SHEET!r} row {row_number}: "
                    f"{barcode}"
                )
            seen_barcodes.add(barcode)
            tasks.append(
                SearchTask(
                    query=name,
                    checkpoint_key=product_checkpoint_key(barcode),
                    kind=SearchKind.PRODUCT,
                    result_policy=ResultPolicy.FIRST,
                    expected_barcode=barcode,
                )
            )

        if not tasks:
            raise ConfigurationError(
                f"Product sheet {PRODUCT_SHEET!r} contains no usable products: {path}"
            )
        return tasks
    finally:
        workbook.close()
